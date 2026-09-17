import os
import fitz
from docx import Document


def extract_pdf_text(pdf_path: str) -> str:
    text = ""

    try:
        pdf = fitz.open(pdf_path)

        for page_number, page in enumerate(pdf, start=1):
            page_text = page.get_text()
            if page_text.strip():
                text += f"\n\n[PAGE {page_number}]\n{page_text}"

        pdf.close()

    except Exception as e:
        print(f"Error extracting PDF {pdf_path}: {e}")

    return text


def extract_txt_text(txt_path: str) -> str:
    try:
        with open(txt_path, "r", encoding="utf-8", errors="ignore") as file:
            return file.read()

    except Exception as e:
        print(f"Error extracting TXT {txt_path}: {e}")
        return ""


def extract_docx_text(docx_path: str) -> str:
    text = ""

    try:
        doc = Document(docx_path)

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"

        # Important: many university documents contain tables.
        # This preserves row relationships instead of losing table data.
        for table in doc.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    text += " | ".join(cells) + "\n"

    except Exception as e:
        print(f"Error extracting DOCX {docx_path}: {e}")

    return text


def get_file_type(file_name: str) -> str:
    extension = os.path.splitext(file_name)[1].lower()

    if extension == ".pdf":
        return "pdf"
    if extension == ".txt":
        return "txt"
    if extension == ".docx":
        return "docx"

    return "unknown"


def extract_all_data(data_folder: str):
    documents = []

    for category in os.listdir(data_folder):
        category_path = os.path.join(data_folder, category)

        if not os.path.isdir(category_path):
            continue

        for root, _, files in os.walk(category_path):
            for file in files:
                file_path = os.path.join(root, file)
                file_type = get_file_type(file)

                if file_type == "pdf":
                    text = extract_pdf_text(file_path)
                elif file_type == "txt":
                    text = extract_txt_text(file_path)
                elif file_type == "docx":
                    text = extract_docx_text(file_path)
                else:
                    continue

                if text.strip():
                    documents.append(
                        {
                            "source": file,
                            "category": category,
                            "file_type": file_type,
                            "file_path": file_path,
                            "text": text,
                        }
                    )

    return documents
